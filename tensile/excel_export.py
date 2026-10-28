"""試験片1本ぶんの「全部入り」Excel ファイルを作る。

寸法のセルを書き換えると、応力・ひずみ・結果・グラフが数式で再計算される。
"""

import io
from datetime import datetime

import numpy as np

EXCEL_POSITIONS = ["上部", "中央", "下部"]


# 「概要」シートのセル（数式から参照する）
CELL_WIDTH = "概要!$B$10"


CELL_THICKNESS = "概要!$B$11"


CELL_GAUGE = "概要!$B$12"


CELL_AREA = "概要!$B$13"


CELL_LOAD_ZERO = "概要!$B$24"


CELL_TIME_ZERO = "概要!$B$25"


CELL_ZERO_DISP = "概要!$B$26"


def to_number_or_blank(value):
    """詳細寸法の文字を数値に直す（空欄や数値でないものは None）。"""
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def recorded_columns(data):
    """補う前の元データ（ファイルに入っていた列と値）を取り出す。"""
    data = data.copy()
    if "Strain_recorded[-]" in data.columns:
        original_columns = list(data.columns[: list(data.columns).index("Strain_recorded[-]")])
        data["Strain[-]"] = data["Strain_recorded[-]"]
        return data[original_columns]
    return data


def excel_value(value):
    """Excel に書ける値に直す（NaN は空欄、numpy の数値は Python の数値）。"""
    if isinstance(value, np.generic):
        value = value.item()
    if isinstance(value, float) and not np.isfinite(value):
        return None
    return value


def write_table(ws, headers, rows):
    """見出し行と、値または数式の行を書く。"""
    ws.append(headers)
    for row in rows:
        ws.append([excel_value(v) for v in row])


def style_table_sheet(ws, header_font, header_fill, widths=None):
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
    ws.freeze_panes = "A2"
    for index, cell in enumerate(ws[1]):
        width = (widths or {}).get(index, max(12, min(30, len(str(cell.value or "")) + 4)))
        ws.column_dimensions[cell.column_letter].width = width


CHART_FONT = "Times New Roman"  # グラフの数字・英字のフォント


def chart_text_properties(size=1000, bold=False):
    """グラフの文字の書式（数字・英字は Times New Roman）。"""
    from openpyxl.chart.text import RichText
    from openpyxl.drawing.text import CharacterProperties, Paragraph, ParagraphProperties
    from openpyxl.drawing.text import Font as DrawingFont

    properties = CharacterProperties(latin=DrawingFont(typeface=CHART_FONT), sz=size, b=bold)
    return RichText(p=[Paragraph(pPr=ParagraphProperties(defRPr=properties), endParaRPr=properties)])


def chart_title(text, size=1100):
    """グラフや軸のタイトル。グラフの中に重ならないように overlay=False にする。"""
    from openpyxl.chart.text import RichText, Text
    from openpyxl.chart.title import Title
    from openpyxl.drawing.text import CharacterProperties, Paragraph, ParagraphProperties, RegularTextRun
    from openpyxl.drawing.text import Font as DrawingFont

    properties = CharacterProperties(latin=DrawingFont(typeface=CHART_FONT), sz=size, b=True)
    rich = RichText(
        p=[Paragraph(pPr=ParagraphProperties(defRPr=properties), r=[RegularTextRun(rPr=properties, t=text)])]
    )
    return Title(tx=Text(rich=rich), overlay=False)


def style_axis(axis, title, position):
    axis.title = chart_title(title, 1000)
    axis.txPr = chart_text_properties(1000)
    axis.delete = False
    axis.axPos = position
    axis.majorTickMark = "out"
    axis.minorTickMark = "none"


def make_scatter_chart(title, x_title, y_title):
    from openpyxl.chart import ScatterChart
    from openpyxl.chart.shapes import GraphicalProperties
    from openpyxl.drawing.line import LineProperties

    # scatterStyle を書かないと、Excel が自動の■マーカーにしてしまう
    chart = ScatterChart(scatterStyle="lineMarker")
    chart.title = chart_title(title, 1200)
    style_axis(chart.x_axis, x_title, "b")
    style_axis(chart.y_axis, y_title, "l")
    # アプリのグラフと同じく、横軸・縦軸とも 0 から始める
    chart.x_axis.scaling.min = 0
    chart.y_axis.scaling.min = 0
    chart.x_axis.majorGridlines = None
    chart.legend.position = "t"
    chart.legend.overlay = False
    chart.legend.txPr = chart_text_properties(1000)
    # グラフの外枠を消す
    chart.graphical_properties = GraphicalProperties(ln=LineProperties(noFill=True))
    chart.roundedCorners = False
    chart.width = 24
    chart.height = 11
    return chart


def add_series(chart, ws_x, x_col, ws_y, y_col, first_row, last_row, title, kind="line", color=None):
    """散布図に系列を足す。kind='line' は線だけ、kind='marker' は × 印だけ、kind='circle' は ● 印だけ。"""
    from openpyxl.chart import Reference, Series
    from openpyxl.chart.shapes import GraphicalProperties

    x_ref = Reference(ws_x, min_col=x_col, min_row=first_row, max_row=last_row)
    y_ref = Reference(ws_y, min_col=y_col, min_row=first_row, max_row=last_row)
    series = Series(y_ref, x_ref, title=title)
    if kind == "line":
        series.marker.symbol = "none"
        series.graphicalProperties.line.width = 19050
        if color:
            series.graphicalProperties.line.solidFill = color
    else:
        series.marker.symbol = "x" if kind == "marker" else "circle"
        series.marker.size = 6 if kind == "marker" else 6
        series.graphicalProperties.line.noFill = True
        marker_properties = GraphicalProperties()
        if kind == "marker":
            marker_properties.noFill = True  # × は線だけで描く（影や塗りなし）
        elif color:
            marker_properties.solidFill = color
        if color:
            marker_properties.line.solidFill = color
        marker_properties.line.width = 9525
        series.marker.graphicalProperties = marker_properties
    series.smooth = False
    chart.series.append(series)
    return series


def make_secondary_chart():
    from openpyxl.chart import ScatterChart

    return ScatterChart(scatterStyle="lineMarker")


def add_secondary(chart, secondary, y_title):
    """2本目の縦軸（右側）を持つグラフを重ねる。"""
    secondary.y_axis.axId = 200
    style_axis(secondary.y_axis, y_title, "r")
    secondary.y_axis.crosses = "max"
    secondary.y_axis.scaling.min = 0
    secondary.y_axis.majorGridlines = None
    chart += secondary


def create_specimen_excel(specimen_name, ss_data, ae_data, ss_result, ae_result, dimension_details):
    """試験片1本ぶんの Excel を作る。

    寸法（黄色のセル）を書き換えると、断面積・応力・ひずみ・最大応力・破断ひずみ・弾性率・
    グラフがすべて自動で変わるように、計算は Excel の数式で書く。
    寸法を変えても、アプリが決めた「ゼロ点の位置（変位）」「直線部分の傾き（荷重/変位）」
    「破断点」は変わらないので、数式の結果はアプリで解析し直した結果と一致する。
    """
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    title_font = Font(bold=True, size=14)
    section_font = Font(bold=True, size=12)
    header_font = Font(bold=True)
    input_font = Font(color="0000FF")
    header_fill = PatternFill("solid", fgColor="DDEBF7")
    input_fill = PatternFill("solid", fgColor="FFFF00")

    width = ss_result["width[mm]"]
    thickness = ss_result["thickness[mm]"]
    gauge = ss_result["gauge_length[mm]"]
    area = ss_result["area[mm2]"]
    zero_disp = ss_result["strain_zero[%]"] / 100 * gauge
    slope_n_per_mm = ss_result["fit_slope"] * area / (gauge / 100)
    fit_indices = np.array(ss_result["fit_indices"], dtype=int)
    fracture_index = int(ss_result["fracture_index"])
    fracture_disp = float(ss_data["Displacement_raw[mm]"].iloc[fracture_index])
    fit_load = ss_data["Load_corrected[N]"].iloc[fit_indices]
    speed = ss_result.get("crosshead_speed[mm/min]")
    n_ss = len(ss_data)
    n_ae = len(ae_data)
    ss_last = n_ss + 1
    ae_last = n_ae + 1

    wb = Workbook()

    # ---------- 概要 ----------
    ws = wb.active
    ws.title = "概要"
    try:
        from zoneinfo import ZoneInfo

        created = datetime.now(ZoneInfo("Asia/Tokyo"))
    except Exception:
        created = datetime.now()
    ws["A1"] = "引張試験・AE解析 結果"
    ws["A1"].font = title_font
    ws["A2"] = (
        "黄色のセル（寸法）を書き換えると、断面積・応力・ひずみ・弾性率・グラフがすべて自動で変わります。"
    )
    ws["A3"] = (
        "寸法は、下の「詳細寸法」に1つでも入力するとその平均（入力したセルだけ）を、空欄なら「直接入力」の値を使います。"
    )
    ws["A4"] = "試験片"
    ws["A5"], ws["B5"] = "試験片名", specimen_name
    ws["A6"], ws["B6"] = "作成日時", created.strftime("%Y-%m-%d %H:%M")
    ws["A7"], ws["B7"], ws["C7"] = (
        "クロスヘッド速度",
        float(f"{speed:.2g}") if speed is not None else None,
        "mm/min",
    )
    ws["A9"], ws["B9"], ws["C9"], ws["D9"] = "寸法", "使う値", "単位", "直接入力"
    ws["E9"] = "← 詳細寸法が空欄のときは、この値を使います"
    ws["A10"], ws["C10"], ws["D10"] = "幅", "mm", width
    ws["A11"], ws["C11"], ws["D11"] = "厚み", "mm", thickness
    ws["A12"], ws["C12"], ws["D12"] = "ゲージ長さ", "mm", gauge
    # 詳細寸法の表（39行目から）の範囲。入力したセルだけの平均、空欄なら直接入力
    ws["B10"] = "=IF(COUNT(B41:D43)>0,AVERAGE(B41:D43),D10)"
    ws["B11"] = "=IF(COUNT(B46:D48)>0,AVERAGE(B46:D48),D11)"
    ws["B12"] = "=IF(COUNT(B51:D51)>0,AVERAGE(B51:D51),D12)"
    ws["A13"], ws["B13"], ws["C13"] = "断面積", "=B10*B11", "mm²"
    ws["A15"] = "結果（自動で計算）"
    ws["A16"], ws["B16"], ws["C16"] = (
        "最大応力（引張強さ）",
        f"=MAX('SS補正後データ'!$I$2:$I${ss_last})",
        "MPa",
    )
    ws["A17"], ws["B17"], ws["C17"] = "破断ひずみ", "=(B28-B26)/B12*100", "%"
    ws["A18"], ws["B18"], ws["C18"] = "破断ひずみ", "=B17/100", "-"
    ws["A19"], ws["B19"], ws["C19"] = "弾性率（見かけ）", "=B27*B12/B13/1000", "GPa"
    ws["A20"], ws["B20"], ws["C20"] = "AEイベント数", ae_result.get("event_count"), "件"
    ws["E19"] = "クロスヘッド変位から求めた見かけの値（試験機・つかみ具のたわみを含むため、実際より小さめ）"
    ws["A22"] = "補正の値（アプリが決めた値。ふつうは変えません）"
    ws["A23"], ws["B23"] = "直線区間の決め方", ss_result.get("fit_mode", "自動")
    ws["A24"], ws["B24"], ws["C24"] = "荷重ゼロ", ss_result["load_zero[N]"], "N"
    ws["A25"], ws["B25"], ws["C25"] = "時間ゼロ", ss_result["time_zero[s]"], "s"
    ws["A26"], ws["B26"], ws["C26"] = "ひずみゼロの位置（変位）", zero_disp, "mm"
    ws["A27"], ws["B27"], ws["C27"] = "直線部分の傾き（荷重÷変位）", slope_n_per_mm, "N/mm"
    ws["A28"], ws["B28"], ws["C28"] = "破断点の変位（補正前）", fracture_disp, "mm"
    ws["A29"], ws["B29"], ws["C29"] = "ひずみゼロ", "=B26/B12*100", "%"
    ws["A30"], ws["B30"], ws["C30"] = "直線区間（荷重）", f"{fit_load.min():.1f}〜{fit_load.max():.1f}", "N"
    ws["A31"], ws["B31"], ws["C31"] = "直線区間の下端（応力）", f"={fit_load.min()}/B13", "MPa"
    ws["A32"], ws["B32"], ws["C32"] = "直線区間の上端（応力）", f"={fit_load.max()}/B13", "MPa"
    ws["A33"], ws["B33"], ws["C33"] = "直線区間の点数", len(fit_indices), "点"
    ws["A34"], ws["B34"] = "直線近似 R²", ss_result["fit_r2"]
    ws["A35"], ws["B35"] = "変位の補完", ss_result.get("disp_gap_message", "") or "なし"
    ws["A36"], ws["B36"] = "注意", ss_result.get("fit_warning", "") or "なし"
    ws["E26"] = "この変位の位置を「ひずみ0」とする（トウ補正）"
    ws["E27"] = "弾性率 ＝ 傾き × ゲージ長さ ÷ 断面積"

    for cell_name in ["A4", "A9", "A15", "A22"]:
        ws[cell_name].font = section_font
    for cell_name in ["B9", "C9", "D9"]:
        ws[cell_name].font = header_font
    for cell_name in ["D10", "D11", "D12"]:
        ws[cell_name].fill = input_fill
        ws[cell_name].font = input_font
    for cell_name in [
        "B10",
        "B11",
        "B12",
        "B13",
        "B16",
        "B17",
        "B18",
        "B19",
        "B26",
        "B27",
        "B28",
        "B29",
        "B31",
        "B32",
    ]:
        ws[cell_name].number_format = "0.000"
    ws["B24"].number_format = "0.000"
    ws["B25"].number_format = "0.000"
    ws["B34"].number_format = "0.00000"
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 10
    ws.column_dimensions["D"].width = 12
    ws.column_dimensions["E"].width = 60

    # 詳細寸法の表（いつも置く。アプリで入力していた値があれば入れておく）
    width_values = [to_number_or_blank(v) for v in dimension_details.get("width_values", [])]
    thickness_values = [to_number_or_blank(v) for v in dimension_details.get("thickness_values", [])]
    gauge_values = [to_number_or_blank(v) for v in dimension_details.get("gauge_values", [])]
    ws.cell(row=39, column=1, value="詳細寸法（入力したセルだけで平均します）").font = section_font
    for title, values, header_row in [("幅 [mm]", width_values, 40), ("厚み [mm]", thickness_values, 45)]:
        ws.cell(row=header_row, column=1, value=title).font = header_font
        for j in range(3):
            ws.cell(row=header_row, column=2 + j, value=f"{j + 1}回目").font = header_font
        for i, position in enumerate(EXCEL_POSITIONS):
            ws.cell(row=header_row + 1 + i, column=1, value=position)
            for j in range(3):
                index = i * 3 + j
                cell = ws.cell(
                    row=header_row + 1 + i, column=2 + j, value=values[index] if index < len(values) else None
                )
                cell.fill = input_fill
                cell.font = input_font
    ws.cell(row=50, column=1, value="ゲージ長さ [mm]").font = header_font
    ws.cell(row=51, column=1, value="測定値")
    for j in range(3):
        ws.cell(row=50, column=2 + j, value=f"{j + 1}回目").font = header_font
        cell = ws.cell(row=51, column=2 + j, value=gauge_values[j] if j < len(gauge_values) else None)
        cell.fill = input_fill
        cell.font = input_font

    # ---------- SS補正後データ（補正後の値は数式） ----------
    ws_ss = wb.create_sheet("SS補正後データ")
    ss_rows = []
    time_raw = ss_data["Time_raw[s]"].to_numpy()
    disp_raw = ss_data["Displacement_raw[mm]"].to_numpy()
    load_raw = ss_data["Load[N]"].to_numpy()
    # グラフ用の列：トウ（直線区間より前）は空欄にし、直線区間の1つ前の行を原点(0, 0)にする。
    # こうすると「原点 → 直線区間 → 以降のデータ」が1つの系列になる
    toe_first_row = int(fit_indices.min()) + 2
    origin_row = toe_first_row - 1 if toe_first_row >= 3 else None
    for i in range(n_ss):
        r = i + 2
        if origin_row is not None and r < origin_row:
            plot_columns = [None, None]
        elif r == origin_row:
            plot_columns = [0, 0]
        else:
            plot_columns = [f"=F{r}", f"=I{r}"]
        ss_rows.append(
            [
                time_raw[i],
                f"=A{r}-{CELL_TIME_ZERO}",
                disp_raw[i],
                f"=C{r}-{CELL_ZERO_DISP}",
                f"=C{r}/{CELL_GAUGE}*100",
                f"=D{r}/{CELL_GAUGE}*100",
                load_raw[i],
                f"=G{r}-{CELL_LOAD_ZERO}",
                f"=H{r}/{CELL_AREA}",
            ]
            + plot_columns
        )
    write_table(
        ws_ss,
        [
            "時間 [s]",
            "時間（補正後） [s]",
            "変位 [mm]",
            "変位（補正後） [mm]",
            "ひずみ（補正前） [%]",
            "ひずみ（補正後） [%]",
            "荷重 [N]",
            "荷重（補正後） [N]",
            "応力 [MPa]",
            "ひずみ（グラフ用） [%]",
            "応力（グラフ用） [MPa]",
        ],
        ss_rows,
    )
    style_table_sheet(ws_ss, header_font, header_fill, widths={i: 16 for i in range(11)})
    ws_ss["M1"] = (
        "「グラフ用」の列：トウの部分を空欄にし、原点(0, 0)から直線区間につなげた曲線（グラフで使用）"
    )

    # ---------- AE補正後データ ----------
    ws_ae = wb.create_sheet("AE補正後データ")
    ae_value_columns = [
        ("Ch.1 Amplitude [mV]", "Ch.1 振幅 [mV]"),
        ("Ch.2 Amplitude [mV]", "Ch.2 振幅 [mV]"),
        ("Ch.1 Energy [V^2/s]", "Ch.1 エネルギー [V²/s]"),
        ("Ch.2 Energy [V^2/s]", "Ch.2 エネルギー [V²/s]"),
        ("Ch.1 Cumulative Energy", "Ch.1 累積エネルギー [V²/s]"),
        ("Ch.2 Cumulative Energy", "Ch.2 累積エネルギー [V²/s]"),
        ("Ch.1 AE Peak Frequency [kHz]", "Ch.1 ピーク周波数 [kHz]"),
        ("Ch.2 AE Peak Frequency [kHz]", "Ch.2 ピーク周波数 [kHz]"),
    ]
    ae_value_columns = [(c, name) for c, name in ae_value_columns if c in ae_data.columns]
    shown = {"Time[ms]", "Strain[-]", "Load[N]", "Strain_recorded[-]"} | {c for c, _ in ae_value_columns}
    original = list(recorded_columns(ae_data).columns)
    other_columns = [c for c in original if c not in shown]
    ae_headers = (
        [
            "時間（AE記録） [s]",
            "時間（SSにそろえた） [s]",
            "変位 [mm]",
            "ひずみ（補正後） [%]",
            "荷重 [N]",
            "応力 [MPa]",
        ]
        + [name for _, name in ae_value_columns]
        + other_columns
    )
    ae_time = ae_data["AE_Time_raw[s]"].to_numpy()
    ae_aligned = ae_data["Time_aligned[s]"].to_numpy()
    ae_disp = (ae_data["Strain[-]"] / 500).to_numpy()
    ae_load = ae_data["Load[N]"].to_numpy() if "Load[N]" in ae_data.columns else np.full(n_ae, np.nan)
    ae_values = [ae_data[c].to_numpy() for c, _ in ae_value_columns]
    ae_others = [ae_data[c].to_numpy() for c in other_columns]
    ae_rows = []
    for i in range(n_ae):
        r = i + 2
        stress_formula = f"=(E{r}-{CELL_LOAD_ZERO})/{CELL_AREA}" if np.isfinite(ae_load[i]) else None
        ae_rows.append(
            [
                ae_time[i],
                ae_aligned[i],
                ae_disp[i],
                f"=(C{r}-{CELL_ZERO_DISP})/{CELL_GAUGE}*100",
                ae_load[i],
                stress_formula,
            ]
            + [values[i] for values in ae_values]
            + [values[i] for values in ae_others]
        )
    write_table(ws_ae, ae_headers, ae_rows)
    style_table_sheet(ws_ae, header_font, header_fill)
    ae_col = {name: ae_headers.index(name) + 1 for name in ae_headers}

    # ---------- 元データ ----------
    ws_ss_raw = wb.create_sheet("SS元データ")
    ss_raw = recorded_columns(ss_data)[["Time[ms]", "Strain[-]", "Load[N]"]]
    write_table(ws_ss_raw, list(ss_raw.columns), ss_raw.itertuples(index=False, name=None))
    style_table_sheet(ws_ss_raw, header_font, header_fill)
    ws_ae_raw = wb.create_sheet("AE元データ")
    ae_raw = recorded_columns(ae_data)
    write_table(ws_ae_raw, list(ae_raw.columns), ae_raw.itertuples(index=False, name=None))
    style_table_sheet(ws_ae_raw, header_font, header_fill)

    # ---------- グラフ（Excel のグラフ。データのセルを参照するので、寸法を変えると変わる） ----------
    ws_graph = wb.create_sheet("グラフ", 1)
    ws_graph["A1"] = (
        "グラフはデータのシートを参照しています。グラフをクリックすると、色や範囲を自由に変えられます。"
    )
    chart_rows = 26  # グラフ1つぶんの行数（グラフの高さ 11 cm ＋ すき間）
    row_height = 15  # 行の高さ [pt] を決めておき、グラフどうしが重ならないようにする
    for row_number in range(1, chart_rows * 12):
        ws_graph.row_dimensions[row_number].height = row_height
    anchor_row = 3

    def place(chart):
        nonlocal anchor_row
        ws_graph.add_chart(chart, f"A{anchor_row}")
        anchor_row += chart_rows

    stress_color = "1F77B4"
    fit_first = int(fit_indices.min()) + 2
    fit_last = int(fit_indices.max()) + 2

    def add_stress(chart):
        """応力-ひずみ曲線（トウ補正済み：原点 → 直線区間 → データ）を1つの系列で足す。"""
        add_series(chart, ws_ss, 10, ws_ss, 11, 2, ss_last, "Stress", color=stress_color)

    chart = make_scatter_chart("Stress-strain curve", "Strain [%]", "Stress [MPa]")
    add_stress(chart)
    chart.legend = None
    place(chart)

    for channel in [1, 2]:
        for value_name, title, y_title, legend_name, kind, color in [
            (
                f"Ch.{channel} 振幅 [mV]",
                f"Stress + AE amplitude (Ch.{channel})",
                "AE amplitude [mV]",
                "AE amplitude",
                "marker",
                "D62728",
            ),
            (
                f"Ch.{channel} 累積エネルギー [V²/s]",
                f"Stress + cumulative AE energy (Ch.{channel})",
                "Cumulative AE energy [V²/s]",
                "Cumulative AE energy",
                "line",
                "2CA02C",
            ),
            (
                f"Ch.{channel} ピーク周波数 [kHz]",
                f"Stress + peak frequency (Ch.{channel})",
                "Peak frequency [kHz]",
                "Peak frequency",
                "marker",
                "D62728",
            ),
        ]:
            if value_name not in ae_col:
                continue
            chart = make_scatter_chart(title, "Strain [%]", "Stress [MPa]")
            add_stress(chart)
            secondary = make_secondary_chart()
            add_series(
                secondary,
                ws_ae,
                ae_col["ひずみ（補正後） [%]"],
                ws_ae,
                ae_col[value_name],
                2,
                ae_last,
                legend_name,
                kind=kind,
                color=color,
            )
            add_secondary(chart, secondary, y_title)
            place(chart)

    # ゼロ点補正の確認（立ち上がり部分）：トウも含めた全データと、直線区間の点
    chart = make_scatter_chart(
        "Zero-point correction (initial region)", "Corrected strain [%]", "Stress [MPa]"
    )
    # このグラフは、ひずみ0より前（トウの部分）も見たいので、自動の範囲にする
    chart.x_axis.scaling.min = None
    chart.y_axis.scaling.min = None
    zoom_last = int(
        np.searchsorted(
            ss_data["Strain_corrected[%]"].to_numpy(),
            ss_data["Strain_corrected[%]"].iloc[fit_indices.max()] * 3,
        )
    )
    zoom_last = min(max(zoom_last, fit_last), ss_last)
    add_series(chart, ws_ss, 6, ws_ss, 9, 2, zoom_last, "Stress (including toe)", color=stress_color)
    add_series(
        chart,
        ws_ss,
        6,
        ws_ss,
        9,
        fit_first,
        fit_last,
        "Points used for the linear fit",
        kind="circle",
        color="FF7F0E",
    )
    place(chart)

    # 開いたときに必ず計算し直す
    wb.calculation.fullCalcOnLoad = True
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()
